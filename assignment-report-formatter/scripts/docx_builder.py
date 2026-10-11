# -*- coding: utf-8 -*-
"""DOCX builder - institutional assignment report (signature format).

Mirrors the PDF design in Word format: double navy page border on the cover
section, structured cover fields table, auto-updating TOC field, navy zebra
tables, shaded code blocks, italic figure captions, running header + page
numbers. Uses fonts that ship with Word (Calibri / Cambria / Consolas).
"""
import os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Emu
from PIL import Image as PILImage

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x5B, 0x64, 0x72)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY_HEX, RULE_HEX, ROWBG_HEX, SUMBG_HEX, CODEBG_HEX = "1F3A5F", "C9CFD8", "F2F5F9", "E4EAF2", "F5F6F8"
CONTENT_W_CM = 16.6


def _rgb(hex6):
    return RGBColor(int(hex6[0:2], 16), int(hex6[2:4], 16), int(hex6[4:6], 16))


def _set_el(parent, tag, **attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(k.replace("_", ":")), v)
    parent.append(el)
    return el


def _cell_shade(cell, hex6):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), hex6)
    tcPr.append(shd)


def _cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for name, spec in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        if spec is None:
            continue
        sz, color = spec
        e = OxmlElement("w:" + name)
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        borders.append(e)
    tcPr.append(borders)


def _table_full_grid(table, color=RULE_HEX, sz=4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + name)
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        borders.append(e)
    tblPr.append(borders)


def _no_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + name)
        e.set(qn("w:val"), "none"); e.set(qn("w:sz"), "0")
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), "auto")
        borders.append(e)
    tblPr.append(borders)


def _page_border(section):
    """Double navy frame around the cover page."""
    sectPr = section._sectPr
    pg = OxmlElement("w:pgBorders")
    pg.set(qn("w:offsetFrom"), "page")
    for name in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + name)
        e.set(qn("w:val"), "double"); e.set(qn("w:sz"), "18")
        e.set(qn("w:space"), "24"); e.set(qn("w:color"), NAVY_HEX)
        pg.append(e)
    sectPr.append(pg)


def _run(par, text, font="Cambria", size=10.5, bold=False, italic=False,
         color=INK, super=False, sub=False):
    r = par.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    if super:
        r.font.superscript = True
    if sub:
        r.font.subscript = True
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:ascii"), font); rf.set(qn("w:hAnsi"), font)
    rf.set(qn("w:eastAsia"), font); rf.set(qn("w:cs"), font)
    return r


TAG_RE = re.compile(r"<(/?)(b|i|super|sub|font)([^>]*)>")


def _apply_inline(par, text, font, size, color, bold, italic):
    """Inline renderer for the markup used in specs: <b> <i> <super> <sub> <font face="Math">.
    Clears any existing runs in the paragraph first, then emits styled runs."""
    for r in list(par.runs):
        r._element.getparent().remove(r._element)
    tokens = re.split(r"(<b>|</b>|<i>|</i>|<super>|</super>|<sub>|</sub>|<font[^>]*>|</font>)", text)
    b = bold; i = italic; sup = False; sub_ = False; stack = []
    cur_font = font
    for tok in tokens:
        if tok in ("", None):
            continue
        if tok == "<b>": b = True; continue
        if tok == "</b>": b = bold; continue
        if tok == "<i>": i = True; continue
        if tok == "</i>": i = italic; continue
        if tok == "<super>": sup = True; continue
        if tok == "</super>": sup = False; continue
        if tok == "<sub>": sub_ = True; continue
        if tok == "</sub>": sub_ = False; continue
        if tok.startswith("<font"):
            face = re.search(r'face\s*=\s*"([^"]+)"', tok)
            if face:
                stack.append(cur_font)
                cur_font = "Cambria" if face.group(1) == "Math" else face.group(1)
            continue
        if tok == "</font>":
            if stack:
                cur_font = stack.pop()
            continue
        _run(par, tok, font=cur_font, size=size, bold=b, italic=i, color=color,
             super=sup, sub=sub_)


def _para(doc_or_cell, text, font="Cambria", size=10.5, bold=False, italic=False,
          color=INK, align=None, before=0, after=8, line=None, keep_next=False):
    p = doc_or_cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if keep_next:
        pf.keep_with_next = True
    if text:
        _apply_inline(p, text, font, size, color, bold, italic)
    return p


def _setup_styles(doc):
    st = doc.styles["Normal"]
    st.font.name = "Cambria"; st.font.size = Pt(10.5); st.font.color.rgb = INK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria")
    pf = st.paragraph_format
    pf.space_after = Pt(8); pf.line_spacing = 1.15
    for name, size, color, before in (("Heading 1", 14.5, NAVY, 18), ("Heading 2", 11.5, INK, 13)):
        h = doc.styles[name]
        h.font.name = "Calibri"; h.font.size = Pt(size); h.font.bold = True
        h.font.color.rgb = color
        h.element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        h.element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(7 if name == "Heading 1" else 5)
        h.paragraph_format.keep_with_next = True


def _add_heading(doc, text, level):
    p = doc.add_paragraph(style="Heading %d" % level)
    font = "Calibri"
    _apply_inline(p, text, font, 14.5 if level == 1 else 11.5,
                  NAVY if level == 1 else INK, True, False)
    return p


def _add_toc_field(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fld = OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"), "begin")
    run._r.append(fld)
    run2 = p.add_run()
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    run2._r.append(instr)
    run3 = p.add_run()
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    run3._r.append(sep)
    _run(p, "Table of contents - right-click and choose \u201cUpdate Field\u201d to fill in page numbers.",
         font="Cambria", size=9.5, italic=True, color=GRAY)
    run4 = p.add_run()
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run4._r.append(end)


def _cover(doc, meta):
    dept = meta.get("dept", "DEPARTMENT OF ARTIFICIAL INTELLIGENCE & MACHINE LEARNING")
    ay = meta.get("ay_line", "AY 2026-27  |  Batch 2025-2029")
    footer = meta.get("footer_line", "Academic Year 2026-27")
    logo = meta.get("logo", os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                         os.pardir, "assets", "sit-logo.png"))

    def cp(text, font, size, color, bold=False, italic=False, before=0, after=0):
        p = _para(doc, text, font=font, size=size, bold=bold, italic=italic, color=color,
                  align=WD_ALIGN_PARAGRAPH.CENTER, before=before, after=after)
        return p

    if logo:
        pw, ph = PILImage.open(logo).size
        lp = _para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, before=24, after=14)
        lp.add_run().add_picture(logo, width=Cm(8.5), height=Emu(int(Cm(8.5) * ph / pw)))
        cp(dept, "Calibri", 10.5, GRAY, bold=True, after=10)
    else:
        cp(dept, "Calibri", 10.5, GRAY, bold=True, before=40, after=10)
    rule = _para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, after=26)
    pPr = rule._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "8")
    b.set(qn("w:space"), "1"); b.set(qn("w:color"), NAVY_HEX)
    pbdr.append(b); pPr.append(pbdr)
    rule.paragraph_format.left_indent = Cm(4.8); rule.paragraph_format.right_indent = Cm(4.8)
    cp(meta.get("course", "Supervised Machine Learning Lab"), "Calibri", 21, INK, bold=True, after=6)
    cp(ay, "Calibri", 11, GRAY, after=52)
    cp(meta["badge"], "Calibri", 12.5, NAVY, bold=True, after=8)
    cp(meta["title"], "Calibri", 28, INK, bold=True, after=10)
    cp(meta["subtitle"], "Calibri", 12, GRAY, italic=True, after=6)
    if meta.get("subtitle2"):
        cp(meta["subtitle2"], "Calibri", 10.5, GRAY, italic=True, after=6)
    _para(doc, "", after=48)

    fields = meta.get("fields") or [
        ("Name", meta["name"]), ("PRN / Roll No.", meta["prn"]),
        ("Assignment", meta.get("assignment", "")),
        ("Course", meta.get("course", "Supervised Machine Learning Lab")),
        ("Date", meta.get("date", "")),
    ]
    ft = doc.add_table(rows=len(fields), cols=2)
    ft.alignment = WD_TABLE_ALIGNMENT.CENTER
    _no_table_borders(ft)
    for i, (k, v) in enumerate(fields):
        c0, c1 = ft.rows[i].cells
        c0.width = Cm(4.6); c1.width = Cm(8.6)
        p0 = c0.paragraphs[0]; p0.paragraph_format.space_before = Pt(6); p0.paragraph_format.space_after = Pt(6)
        _run(p0, k, font="Calibri", size=10.5, bold=True, color=NAVY)
        p1 = c1.paragraphs[0]; p1.paragraph_format.space_before = Pt(6); p1.paragraph_format.space_after = Pt(6)
        _run(p1, v, font="Cambria", size=11, color=INK)
        top = (8, NAVY_HEX) if i == 0 else (3, RULE_HEX)
        bottom = (8, NAVY_HEX) if i == len(fields) - 1 else (3, RULE_HEX)
        _cell_borders(c0, top=top, bottom=bottom)
        _cell_borders(c1, top=top, bottom=bottom)
    _para(doc, footer, font="Calibri", size=9.5, color=GRAY,
          align=WD_ALIGN_PARAGRAPH.CENTER, before=44)


def _header_footer(section, left, right):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_W_CM), WD_TAB_ALIGNMENT.RIGHT)
    _run(hp, left, font="Calibri", size=8.5, color=GRAY)
    _run(hp, "\t" + right, font="Calibri", size=8.5, color=GRAY)
    pPr = hp._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4")
    b.set(qn("w:space"), "4"); b.set(qn("w:color"), RULE_HEX)
    pbdr.append(b); pPr.append(pbdr)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = fp._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:top")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4")
    b.set(qn("w:space"), "4"); b.set(qn("w:color"), RULE_HEX)
    pbdr.append(b); pPr.append(pbdr)
    _run(fp, "Page ", font="Calibri", size=8.5, color=GRAY)
    r = fp.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    r._r.append(f1)
    r2 = fp.add_run()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "PAGE"
    r2._r.append(it)
    r3 = fp.add_run()
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r3._r.append(f2)
    for rr in (r, r2, r3):
        rr.font.name = "Calibri"; rr.font.size = Pt(8.5); rr.font.color.rgb = GRAY


def _data_table(doc, header, rows, widths_cm, aligns, summary_rows=None, math_cols=None):
    math_cols = math_cols or []
    summary_rows = set(summary_rows or [])
    t = doc.add_table(rows=len(rows) + 1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _table_full_grid(t)
    for ci, h in enumerate(header):
        c = t.rows[0].cells[ci]
        c.width = Cm(widths_cm[ci])
        _cell_shade(c, NAVY_HEX)
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
        _apply_inline(p, h, "Calibri", 9, WHITE, True, False)
    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            c = t.rows[ri + 1].cells[ci]
            c.width = Cm(widths_cm[ci])
            if ri in summary_rows:
                _cell_shade(c, SUMBG_HEX)
            elif len(summary_rows) == 0 and ri % 2 == 1:
                _cell_shade(c, ROWBG_HEX)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if aligns[ci] == "c" else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
            if ri in summary_rows or aligns[ci] == "b":
                _apply_inline(p, str(cell_text), "Cambria", 9.5, INK, True, False)
            elif ci in math_cols:
                _apply_inline(p, str(cell_text), "Cambria", 9.5, INK, False, False)
            else:
                _apply_inline(p, str(cell_text), "Cambria", 9.5, INK, False, False)
    return t


def _code_block(doc, text):
    lines = text.rstrip("\n").split("\n")
    t = doc.add_table(rows=len(lines), cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + name)
        e.set(qn("w:val"), "none"); e.set(qn("w:sz"), "0")
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), "auto")
        borders.append(e)
    e = OxmlElement("w:left")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "18")
    e.set(qn("w:space"), "0"); e.set(qn("w:color"), NAVY_HEX)
    borders.append(e)
    tblPr.append(borders)
    for i, ln in enumerate(lines):
        c = t.rows[i].cells[0]
        c.width = Cm(CONTENT_W_CM)
        _cell_shade(c, CODEBG_HEX)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        _run(p, ln if ln.strip() else " ", font="Consolas", size=7.9, color=INK)
    return t


def build_docx(spec, out_path):
    meta = spec["meta"]
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2.3); sec.bottom_margin = Cm(2.0)
    _setup_styles(doc)
    _page_border(sec)
    _cover(doc, meta)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    # remove inherited cover page border from the body section
    _pg = body._sectPr.find(qn("w:pgBorders"))
    if _pg is not None:
        body._sectPr.remove(_pg)
    body.page_width = Cm(21.0); body.page_height = Cm(29.7)
    body.left_margin = Cm(2.2); body.right_margin = Cm(2.2)
    body.top_margin = Cm(2.3); body.bottom_margin = Cm(2.0)
    hl = meta.get("header_left", "%s  |  AY 2026-27" % meta.get("course", ""))
    _header_footer(body, hl, meta.get("header_right", meta["title"]))

    _para(doc, "Table of Contents", font="Calibri", size=16, bold=True, color=NAVY, after=6)
    hr = _para(doc, "", after=14)
    pPr = hr._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "8")
    b.set(qn("w:space"), "1"); b.set(qn("w:color"), NAVY_HEX)
    pbdr.append(b); pPr.append(pbdr)
    _add_toc_field(doc)
    doc.add_page_break()

    for blk in spec["sections"]:
        ty = blk["type"]
        if ty == "h1":
            _add_heading(doc, blk["text"], 1)
        elif ty == "h2":
            _add_heading(doc, blk["text"], 2)
        elif ty == "p":
            _para(doc, blk["text"], align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif ty == "lead":
            _para(doc, blk["text"], bold=True, before=4, after=4)
        elif ty == "bullets":
            for it in blk["items"]:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(3)
                _apply_inline(p, it, "Cambria", 10.5, INK, False, False)
            _para(doc, "", after=4)
        elif ty == "numbered":
            for it in blk["items"]:
                p = doc.add_paragraph(style="List Number")
                p.paragraph_format.space_after = Pt(3)
                _apply_inline(p, it, "Cambria", 10.5, INK, False, False)
            _para(doc, "", after=4)
        elif ty == "formula":
            _para(doc, blk["text"], align=WD_ALIGN_PARAGRAPH.CENTER, before=3, after=9)
        elif ty == "spacer":
            _para(doc, "", after=blk.get("cm", 0.4) * 28.35)
        elif ty == "code":
            _code_block(doc, blk["text"])
            _para(doc, "", after=10)
        elif ty == "figure":
            path = blk["path"]; wcm = blk.get("width_cm", 16.0)
            pw, ph = PILImage.open(path).size
            w = Cm(min(wcm, CONTENT_W_CM)); h = Emu(int(w * ph / pw))
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(path, width=w, height=h)
            cp = _para(doc, blk["caption"], font="Calibri", size=9, italic=True, color=GRAY,
                       align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
        elif ty == "table":
            widths = blk.get("widths") or [CONTENT_W_CM / len(blk["header"])] * len(blk["header"])
            _data_table(doc, blk["header"], blk["rows"], widths,
                        blk.get("align", ["l"] * len(blk["header"])),
                        blk.get("summary_rows"), blk.get("math_cols"))
            _para(doc, "", after=10)
        elif ty == "kv_table":
            w1 = blk.get("w1", 6.0); w2 = blk.get("w2", 10.6)
            t = doc.add_table(rows=len(blk["rows"]), cols=2)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            _table_full_grid(t)
            for ri, (k, v) in enumerate(blk["rows"]):
                c0, c1 = t.rows[ri].cells
                c0.width = Cm(w1); c1.width = Cm(w2)
                _cell_shade(c0, ROWBG_HEX if ri % 2 == 0 else "FFFFFF")
                _cell_shade(c1, ROWBG_HEX if ri % 2 == 0 else "FFFFFF")
                p0 = c0.paragraphs[0]; p1 = c1.paragraphs[0]
                p0.paragraph_format.space_before = Pt(4); p0.paragraph_format.space_after = Pt(4)
                p1.paragraph_format.space_before = Pt(4); p1.paragraph_format.space_after = Pt(4)
                _apply_inline(p0, k, "Calibri", 9, NAVY, True, False)
                _apply_inline(p1, str(v), "Cambria", 9.5, INK, True, False)
            _para(doc, "", after=10)
        else:
            raise ValueError("unknown block type: %s" % ty)

    doc.save(out_path)
    return out_path
